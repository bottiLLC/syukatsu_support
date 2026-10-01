# Copyright (C) 2026 合同会社ぼっち (bottiLLC)
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

"""MarkdownテキストをMicrosoft Word (.docx) 文書へ変換・エクスポートするモジュール。

見出し (H1-H4)、箇条書き、番号付きリスト、強調（太字）、引用、テーブル、
および通常段落を解析し、Word文書スタイルへマッピングして保存します。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Final

import docx
from docx.document import Document as DocxDocument
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

_HEADING_PATTERN: Final[re.Pattern[str]] = re.compile(r"^(#{1,4})\s+(.+)$")
_BULLET_PATTERN: Final[re.Pattern[str]] = re.compile(r"^[-*・]\s+(.+)$")
_NUMBERED_PATTERN: Final[re.Pattern[str]] = re.compile(r"^\d+[.)]\s+(.+)$")
_QUOTE_PATTERN: Final[re.Pattern[str]] = re.compile(r"^>\s*(.+)$")
_BOLD_SPLIT_PATTERN: Final[re.Pattern[str]] = re.compile(r"(\*\*.*?\*\*)")


def _add_formatted_runs(paragraph: Paragraph, text: str) -> None:
    """段落に対してインライン装飾（太字）を解析してRunを追加します。

    Args:
        paragraph: 対象のWord段落オブジェクト
        text: 装飾を含むテキスト文字列
    """
    parts = _BOLD_SPLIT_PATTERN.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def _parse_table_row(line: str) -> list[str]:
    """Markdownテーブルの行文字列をセル単位のリストに分割します。"""
    trimmed = line.strip()
    if trimmed.startswith("|"):
        trimmed = trimmed[1:]
    if trimmed.endswith("|"):
        trimmed = trimmed[:-1]
    return [cell.strip() for cell in trimmed.split("|")]


def _is_table_separator(line: str) -> bool:
    """Markdownテーブルの区切り行（|---|---|等）であるかを判定します。"""
    cells = _parse_table_row(line)
    return len(cells) > 0 and all(re.match(r"^:?-+:?$", cell) for cell in cells if cell)


def _render_table(doc: DocxDocument, table_lines: list[str]) -> None:
    """Markdownテーブル行のリストをWordテーブルとしてドキュメントに挿入します。"""
    rows_data: list[list[str]] = []
    for line in table_lines:
        if _is_table_separator(line):
            continue
        rows_data.append(_parse_table_row(line))

    if not rows_data:
        return

    max_cols = max(len(row) for row in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=max_cols)
    table.style = "Table Grid"

    for r_idx, row in enumerate(rows_data):
        for c_idx in range(max_cols):
            cell_text = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = cell_text
            # ヘッダー行は太字に設定
            if r_idx == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

    # テーブルの後に空行を追加
    doc.add_paragraph()


def export_markdown_to_docx(markdown_text: str, output_path: str | Path) -> Path:
    """Markdown形式のテキストをWord (.docx) ファイルとして出力します。

    Args:
        markdown_text: Markdown形式のレポートテキスト
        output_path: 出力先ファイルパス

    Returns:
        Path: 出力されたdocxファイルの絶対パス
    """
    target_path = Path(output_path).resolve()
    target_path.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # ドキュメントの基本フォント設定
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Segoe UI"
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)

    lines = markdown_text.splitlines()
    table_buffer: list[str] = []

    def flush_table() -> None:
        if table_buffer:
            _render_table(doc, table_buffer)
            table_buffer.clear()

    for line in lines:
        stripped = line.strip()

        # テーブル行の判定
        if stripped.startswith("|") and stripped.endswith("|"):
            table_buffer.append(stripped)
            continue

        # テーブル終了検知
        flush_table()

        if not stripped:
            continue

        # 見出し判定 (#, ##, ###, ####)
        heading_match = _HEADING_PATTERN.match(stripped)
        if heading_match:
            hashes, heading_text = heading_match.groups()
            level = min(len(hashes), 4)
            p = doc.add_heading(level=level)
            _add_formatted_runs(p, heading_text)
            continue

        # 箇条書きリスト判定 (-, *, ・)
        bullet_match = _BULLET_PATTERN.match(stripped)
        if bullet_match:
            bullet_text = bullet_match.group(1)
            p = doc.add_paragraph(style="List Bullet")
            _add_formatted_runs(p, bullet_text)
            continue

        # 番号付きリスト判定 (1. 2.)
        numbered_match = _NUMBERED_PATTERN.match(stripped)
        if numbered_match:
            numbered_text = numbered_match.group(1)
            p = doc.add_paragraph(style="List Number")
            _add_formatted_runs(p, numbered_text)
            continue

        # 引用判定 (>)
        quote_match = _QUOTE_PATTERN.match(stripped)
        if quote_match:
            quote_text = quote_match.group(1)
            p = doc.add_paragraph(style="Quote")
            _add_formatted_runs(p, quote_text)
            continue

        # 通常段落
        p = doc.add_paragraph()
        _add_formatted_runs(p, stripped)

    flush_table()

    doc.save(str(target_path))
    return target_path
