# Copyright (C) 2026 合同会社ぼっち (bottiLLC)
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

"""DOCXエクスポート機能のユニットテストモジュール。"""

from pathlib import Path

import docx

from src.core.docx_exporter import export_markdown_to_docx


def test_export_markdown_to_docx_comprehensive(tmp_path: Path) -> None:
    """Markdownの各種要素（見出し、箇条書き、番号、太字、引用、テーブル）が正しくdocx化されることを検証します。"""
    sample_md = """# 大見出し レポート
## 中見出し セクション
### 小見出し サブセクション

これは通常の段落であり、**重要なキーワード**が含まれます。

- 箇条書きアイテム1
* 箇条書きアイテム2
・ 箇条書きアイテム3

1. 番号付きアイテム1
2. 番号付きアイテム2

> これは重要な引用ブロックです。

| 指標名 | 前期実績 | 当期実績 |
| :--- | :---: | ---: |
| 売上高 | 100億円 | 120億円 |
| 営業利益 | 10億円 | 15億円 |
"""
    output_file = tmp_path / "test_report.docx"
    result_path = export_markdown_to_docx(sample_md, output_file)

    assert result_path.exists()
    assert result_path == output_file.resolve()

    # 出力されたdocxを読み込み、構造を検証
    doc = docx.Document(str(result_path))

    # 見出しと段落テキストの確認
    texts = [p.text for p in doc.paragraphs if p.text]
    assert "大見出し レポート" in texts
    assert "中見出し セクション" in texts
    assert "小見出し サブセクション" in texts
    assert "これは通常の段落であり、重要なキーワードが含まれます。" in texts
    assert "箇条書きアイテム1" in texts
    assert "番号付きアイテム1" in texts
    assert "これは重要な引用ブロックです。" in texts

    # テーブルの検証
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert table.cell(0, 0).text == "指標名"
    assert table.cell(1, 0).text == "売上高"
    assert table.cell(1, 1).text == "100億円"
    assert table.cell(2, 2).text == "15億円"


def test_export_markdown_edge_cases(tmp_path: Path) -> None:
    """太字開始、不揃いなカラム数のテーブル、空行等のエッジケースを検証します。"""
    sample_md = """**先頭太字**のテキストです。

| A | B | C |
|---|---|---|
| 1 | 2 |
"""
    output_file = tmp_path / "edge_case.docx"
    result_path = export_markdown_to_docx(sample_md, output_file)
    assert result_path.exists()

    doc = docx.Document(str(result_path))
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 0).text == "1"
    assert doc.tables[0].cell(1, 2).text == ""
